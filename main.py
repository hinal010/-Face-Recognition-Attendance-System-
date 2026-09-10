import os
import base64
from datetime import datetime, date, timedelta
from typing import Optional
from io import BytesIO

import face_recognition
import numpy as np
import plotly.graph_objects as go

from fastapi import (
    FastAPI,
    Request,
    Depends,
    Form,
    HTTPException,
    Query
)

from fastapi.responses import (
    HTMLResponse,
    StreamingResponse
)

from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from sqlalchemy.orm import Session

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from database import Base, engine, get_db
from models import Person, FaceEncoding, Attendance, Holiday
from face_utils import get_face_encoding, string_to_encoding

from zoneinfo import ZoneInfo


# ============================================================
# DATABASE
# ============================================================

Base.metadata.create_all(bind=engine)


# ============================================================
# FASTAPI
# ============================================================

app = FastAPI()


# ============================================================
# STATIC / TEMPLATES
# ============================================================

app.mount(
    "/static",
    StaticFiles(directory="static"),
    name="static"
)

templates = Jinja2Templates(directory="templates")


# ============================================================
# DIRECTORIES / SETTINGS
# ============================================================

UPLOAD_DIR = "uploads"

os.makedirs(UPLOAD_DIR, exist_ok=True)

MIN_OUT_TIME = timedelta(minutes=10)

IST = ZoneInfo("Asia/Kolkata")


# ============================================================
# PAGE ROUTES
# ============================================================

@app.get("/", response_class=HTMLResponse)
def dashboard(
    request: Request,
    db: Session = Depends(get_db)
):

    total_registered = db.query(Person).count()

    today = datetime.now(IST).date()

    today_present = db.query(Attendance).filter(
        Attendance.date == today
    ).count()

    today_absent = total_registered - today_present

    recent_attendance = (
        db.query(Attendance)
        .order_by(Attendance.id.desc())
        .limit(5)
        .all()
    )

    persons = (
        db.query(Person)
        .order_by(Person.full_name.asc())
        .all()
    )

    return templates.TemplateResponse(
        request,
        "dashboard.html",
        {
            "total_registered": total_registered,
            "today_present": today_present,
            "today_absent": today_absent,
            "recent_attendance": recent_attendance,
            "persons": persons
        }
    )


@app.get("/register-page", response_class=HTMLResponse)
def register_page(request: Request):

    return templates.TemplateResponse(
        request,
        "register.html"
    )


@app.get("/attendance-page", response_class=HTMLResponse)
def attendance_page(request: Request):

    return templates.TemplateResponse(
        request,
        "attendance.html"
    )


@app.get("/persons-page", response_class=HTMLResponse)
def persons_page(
    request: Request,
    db: Session = Depends(get_db)
):

    persons = (
        db.query(Person)
        .order_by(Person.id.desc())
        .all()
    )

    return templates.TemplateResponse(
        request,
        "persons.html",
        {
            "persons": persons
        }
    )


@app.get("/holidays-page", response_class=HTMLResponse)
def holidays_page(
    request: Request,
    db: Session = Depends(get_db)
):

    holidays = (
        db.query(Holiday)
        .order_by(Holiday.start_date.desc())
        .all()
    )

    return templates.TemplateResponse(
        request,
        "holidays.html",
        {
            "holidays": holidays
        }
    )


# ============================================================
# FACE DUPLICATE CHECK
# ============================================================

def check_duplicate_face(
    new_encodings,
    db: Session,
    tolerance=0.45
):

    all_encodings = db.query(FaceEncoding).all()

    for new_encoding in new_encodings:

        for stored in all_encodings:

            stored_encoding = string_to_encoding(
                stored.encoding
            )

            distance = face_recognition.face_distance(
                [stored_encoding],
                new_encoding
            )[0]

            if distance <= tolerance:

                return stored.person, distance

    return None, None


# ============================================================
# REGISTER PERSON
# ============================================================

@app.post("/register")
def register_person(
    person_id: str = Form(...),
    full_name: str = Form(...),
    email: str = Form(...),
    contact: str = Form(...),
    college: str = Form(None),
    semester: str = Form(None),
    department: str = Form(...),

    straight_image: str = Form(...),
    left_image: str = Form(...),
    right_image: str = Form(...),

    db: Session = Depends(get_db)
):

    existing_person = (
        db.query(Person)
        .filter(Person.person_id == person_id)
        .first()
    )

    if existing_person:

        raise HTTPException(
            status_code=400,
            detail="Person ID already exists"
        )

    images = {
        "straight": straight_image,
        "left": left_image,
        "right": right_image
    }

    image_paths = {}
    new_encodings = {}

    temp_encoding_arrays = []

    try:

        for pose_name, image_data in images.items():

            image_data = image_data.split(",")[1]

            image_bytes = base64.b64decode(
                image_data
            )

            image_path = os.path.join(
                UPLOAD_DIR,
                f"{person_id}_{pose_name}.jpg"
            )

            with open(image_path, "wb") as file:

                file.write(image_bytes)

            encoding_string, message = get_face_encoding(
                image_path
            )

            if message != "success":

                for path in image_paths.values():

                    if os.path.exists(path):
                        os.remove(path)

                if os.path.exists(image_path):
                    os.remove(image_path)

                raise HTTPException(
                    status_code=400,
                    detail=f"{pose_name} image error: {message}"
                )

            encoding_array = string_to_encoding(
                encoding_string
            )

            image_paths[pose_name] = image_path

            new_encodings[pose_name] = encoding_string

            temp_encoding_arrays.append(
                encoding_array
            )

        duplicate_person, distance = check_duplicate_face(
            temp_encoding_arrays,
            db
        )

        if duplicate_person:

            for path in image_paths.values():

                if os.path.exists(path):
                    os.remove(path)

            raise HTTPException(
                status_code=400,
                detail=(
                    f"This face is already registered as "
                    f"{duplicate_person.full_name} "
                    f"({duplicate_person.person_id})"
                )
            )

        new_person = Person(
            person_id=person_id,
            full_name=full_name,
            email=email,
            contact=contact,
            college=college,
            semester=semester,
            department=department
        )

        db.add(new_person)

        db.commit()

        db.refresh(new_person)

        for pose_name, encoding_string in new_encodings.items():

            face_encoding = FaceEncoding(
                person_id_fk=new_person.id,
                pose_name=pose_name,
                image_path=image_paths[pose_name],
                encoding=encoding_string
            )

            db.add(face_encoding)

        db.commit()

    except HTTPException:
        raise

    except Exception as e:

        for path in image_paths.values():

            if os.path.exists(path):
                os.remove(path)

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )

    return {
        "message": "Registration successful",
        "person_id": person_id,
        "full_name": full_name,
        "department": department
    }


# ============================================================
# MARK ATTENDANCE
# ============================================================

@app.post("/mark-attendance")
def mark_attendance(
    image_data: str = Form(...),
    db: Session = Depends(get_db)
):

    temp_path = os.path.join(
        UPLOAD_DIR,
        "temp_attendance.jpg"
    )

    try:

        image_data = image_data.split(",")[1]

        image_bytes = base64.b64decode(
            image_data
        )

        with open(temp_path, "wb") as file:

            file.write(image_bytes)

        image = face_recognition.load_image_file(
            temp_path
        )

        face_locations = face_recognition.face_locations(
            image
        )

        if len(face_locations) == 0:

            return {
                "status": "waiting",
                "message": "No face detected"
            }

        if len(face_locations) > 1:

            return {
                "status": "multiple",
                "message": "Only one face should be visible"
            }

        current_encodings = face_recognition.face_encodings(
            image,
            face_locations
        )

        if len(current_encodings) == 0:

            return {
                "status": "error",
                "message": "Face encoding failed"
            }

        current_encoding = current_encodings[0]

        all_encodings = (
            db.query(FaceEncoding)
            .all()
        )

        if not all_encodings:

            return {
                "status": "empty",
                "message": "No registered persons found"
            }

        best_match_person = None

        best_distance = 1.0

        for stored in all_encodings:

            stored_encoding = string_to_encoding(
                stored.encoding
            )

            distance = face_recognition.face_distance(
                [stored_encoding],
                current_encoding
            )[0]

            if distance < best_distance:

                best_distance = distance

                best_match_person = stored.person

        if best_distance > 0.45:

            return {
                "status": "unknown",
                "message": "Unknown person"
            }

        ist_now = datetime.now(IST)

        today = ist_now.date()

        current_time = ist_now.time()

        existing_attendance = (
            db.query(Attendance)
            .filter(
                Attendance.person_id_fk == best_match_person.id,
                Attendance.date == today
            )
            .first()
        )

        # ----------------------------------------------------
        # FIRST ENTRY
        # ----------------------------------------------------

        if not existing_attendance:

            attendance = Attendance(
                person_id_fk=best_match_person.id,
                date=today,
                in_time=current_time,
                out_time=None,
                status="Present"
            )

            db.add(attendance)

            db.commit()

            return {
                "status": "in_marked",
                "message": "In Time marked successfully",
                "person_id": best_match_person.person_id,
                "full_name": best_match_person.full_name,
                "department": best_match_person.department,
                "date": str(today),
                "in_time": str(current_time),
                "out_time": "-"
            }

        # ----------------------------------------------------
        # OUT TIME
        # ----------------------------------------------------

        if existing_attendance.out_time is None:

            in_datetime = datetime.combine(
                today,
                existing_attendance.in_time
            )

            current_datetime = datetime.combine(
                today,
                current_time
            )

            time_difference = (
                current_datetime -
                in_datetime
            )

            if time_difference >= MIN_OUT_TIME:

                existing_attendance.out_time = current_time

                db.commit()

                return {
                    "status": "out_marked",
                    "message": "Out Time marked successfully",
                    "person_id": best_match_person.person_id,
                    "full_name": best_match_person.full_name,
                    "department": best_match_person.department,
                    "date": str(existing_attendance.date),
                    "in_time": str(existing_attendance.in_time),
                    "out_time": str(
                        existing_attendance.out_time
                    )
                }

            return {
                "status": "already_in",
                "message": (
                    "In Time already marked. "
                    "Out Time allowed after 10 min."
                ),
                "person_id": best_match_person.person_id,
                "full_name": best_match_person.full_name,
                "department": best_match_person.department,
                "date": str(existing_attendance.date),
                "in_time": str(existing_attendance.in_time),
                "out_time": "-"
            }

        # ----------------------------------------------------
        # COMPLETED
        # ----------------------------------------------------

        return {
            "status": "completed",
            "message": (
                "In Time and Out Time "
                "already marked today"
            ),
            "person_id": best_match_person.person_id,
            "full_name": best_match_person.full_name,
            "department": best_match_person.department,
            "date": str(existing_attendance.date),
            "in_time": existing_attendance.in_time.strftime(
                "%H:%M:%S"
            ),
            "out_time": existing_attendance.out_time.strftime(
                "%H:%M:%S"
            )
        }

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )

    finally:

        if os.path.exists(temp_path):
            os.remove(temp_path)


# ============================================================
# PERSON ATTENDANCE PAGE
# ============================================================

@app.get(
    "/person-attendance/{person_id}",
    response_class=HTMLResponse
)
def person_attendance(
    person_id: str,
    request: Request,
    db: Session = Depends(get_db)
):

    person = (
        db.query(Person)
        .filter(Person.person_id == person_id)
        .first()
    )

    # Support old numeric database IDs as well.
    if not person and person_id.isdigit():
        person = (
            db.query(Person)
            .filter(Person.id == int(person_id))
            .first()
        )

    if not person:

        raise HTTPException(
            status_code=404,
            detail="Person not found"
        )

    attendances = (
        db.query(Attendance)
        .filter(
            Attendance.person_id_fk == person.id
        )
        .all()
    )

    start_date = person.created_at.date()

    end_date = datetime.now(IST).date()

    report = []

    current_date = start_date

    holidays = db.query(Holiday).all()

    while current_date <= end_date:

        attendance = next(
            (
                a
                for a in attendances
                if a.date == current_date
            ),
            None
        )

        holiday_found = None

        for holiday in holidays:

            if (
                holiday.start_date
                <= current_date
                <= holiday.end_date
            ):

                holiday_found = holiday

                break

        if attendance:

            status = "Present"

            in_time = attendance.in_time

            out_time = attendance.out_time

        elif current_date.weekday() == 6:

            status = "Sunday Holiday"

            in_time = None

            out_time = None

        elif holiday_found:

            status = holiday_found.holiday_name

            in_time = None

            out_time = None

        else:

            status = "Absent"

            in_time = None

            out_time = None

        report.append(
            {
                "date": current_date,
                "day": current_date.strftime("%A"),
                "status": status,
                "in_time": in_time,
                "out_time": out_time
            }
        )

        current_date += timedelta(days=1)

    return templates.TemplateResponse(
        request,
        "person_attendance.html",
        {
            "person": person,
            "report": report
        }
    )


# ============================================================
# DELETE PERSON
# ============================================================

@app.post("/delete-person/{person_id}")
def delete_person(
    person_id: int,
    db: Session = Depends(get_db)
):

    person = (
        db.query(Person)
        .filter(Person.id == person_id)
        .first()
    )

    if not person:

        raise HTTPException(
            status_code=404,
            detail="Person not found"
        )

    db.delete(person)

    db.commit()

    return {
        "message": "Person deleted successfully"
    }


# ============================================================
# GRAPH DATA
# ============================================================

@app.get("/get-graph-data")
def get_graph_data(
    db: Session = Depends(get_db)
):

    today = datetime.now(IST).date()

    labels = []

    present_values = []

    for i in range(6, -1, -1):

        day = today - timedelta(days=i)

        present_count = (
            db.query(Attendance)
            .filter(
                Attendance.date == day
            )
            .count()
        )

        labels.append(
            day.strftime("%d %b")
        )

        present_values.append(
            present_count
        )

    return {
        "labels": labels,
        "values": present_values
    }


# ============================================================
# ADD HOLIDAY
# ============================================================

@app.post("/add-holiday")
def add_holiday(
    holiday_name: str = Form(...),
    holiday_type: str = Form(...),
    start_date: str = Form(...),
    end_date: str = Form(...),
    db: Session = Depends(get_db)
):

    start_date_obj = datetime.strptime(
        start_date,
        "%Y-%m-%d"
    ).date()

    end_date_obj = datetime.strptime(
        end_date,
        "%Y-%m-%d"
    ).date()

    if end_date_obj < start_date_obj:

        raise HTTPException(
            status_code=400,
            detail="End date cannot be before start date"
        )

    holiday = Holiday(
        holiday_name=holiday_name,
        holiday_type=holiday_type,
        start_date=start_date_obj,
        end_date=end_date_obj
    )

    db.add(holiday)

    db.commit()

    return {
        "message": "Holiday added successfully"
    }


# ============================================================
# DELETE HOLIDAY
# ============================================================

@app.post("/delete-holiday/{holiday_id}")
def delete_holiday(
    holiday_id: int,
    db: Session = Depends(get_db)
):

    holiday = (
        db.query(Holiday)
        .filter(Holiday.id == holiday_id)
        .first()
    )

    if not holiday:

        raise HTTPException(
            status_code=404,
            detail="Holiday not found"
        )

    db.delete(holiday)

    db.commit()

    return {
        "message": "Holiday deleted successfully"
    }


# ============================================================
# FACE POSITION CHECK
# ============================================================

@app.post("/check-face-position")
def check_face_position(
    image_data: str = Form(...)
):

    temp_path = os.path.join(
        UPLOAD_DIR,
        "temp_check.jpg"
    )

    try:

        image_data = image_data.split(",")[1]

        image_bytes = base64.b64decode(
            image_data
        )

        with open(temp_path, "wb") as file:

            file.write(image_bytes)

        image = face_recognition.load_image_file(
            temp_path
        )

        face_locations = face_recognition.face_locations(
            image
        )

        height, width, _ = image.shape

        if len(face_locations) == 0:

            return {
                "ready": False,
                "message": "No face detected"
            }

        if len(face_locations) > 1:

            return {
                "ready": False,
                "message": "Only one face allowed"
            }

        top, right, bottom, left = face_locations[0]

        face_width = right - left

        if face_width < width * 0.15:

            return {
                "ready": False,
                "message": "Move closer"
            }

        if face_width > width * 0.60:

            return {
                "ready": False,
                "message": "Move back"
            }

        return {
            "ready": True,
            "message": "Face detected",
            "box": {
                "top": top,
                "right": right,
                "bottom": bottom,
                "left": left,
                "image_width": width,
                "image_height": height
            }
        }

    finally:

        if os.path.exists(temp_path):

            os.remove(temp_path)


# ============================================================
# NEXT PERSON ID
# ============================================================

@app.get("/next-person-id")
def next_person_id(
    db: Session = Depends(get_db)
):

    # Generate new IDs from existing P### IDs, not database primary keys.
    # Old IDs such as 001 remain unchanged.
    import re

    persons = db.query(Person.person_id).all()

    max_number = 0

    for (raw_id,) in persons:
        value = (raw_id or "").strip().upper()
        match = re.fullmatch(r"P(\d+)", value)

        if match:
            max_number = max(
                max_number,
                int(match.group(1))
            )

    next_number = max_number + 1

    while db.query(Person).filter(
        Person.person_id == f"P{next_number:03d}"
    ).first():
        next_number += 1

    return {
        "person_id": f"P{next_number:03d}"
    }


# ============================================================
# ============================================================
# EXCEL REPORT FUNCTIONS
# ============================================================
# ============================================================


# ============================================================
# EXCEL STYLE
# ============================================================

def style_excel_sheet(ws):

    # Header style
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.border = thin_border

    # First row header
    for cell in ws[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


def auto_width_columns(ws):

    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                length = len(
                    str(cell.value)
                )

                if length > max_length:

                    max_length = length

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            40
        )


# ============================================================
# GET HOLIDAY STATUS
# ============================================================

def get_date_status(
    current_date,
    attendance,
    holidays
):

    if attendance:

        return (
            "Present",
            attendance.in_time,
            attendance.out_time
        )

    # Sunday
    if current_date.weekday() == 6:

        return (
            "Sunday Holiday",
            None,
            None
        )

    # Other holiday
    for holiday in holidays:

        if (
            holiday.start_date
            <= current_date
            <= holiday.end_date
        ):

            return (
                holiday.holiday_name,
                None,
                None
            )

    # No attendance and no holiday
    return (
        "Absent",
        None,
        None
    )


# ============================================================
# PERSON-WISE REPORT DATA
# ============================================================

def build_person_report_data(
    person,
    db,
    from_date,
    to_date
):

    attendances = (
        db.query(Attendance)
        .filter(
            Attendance.person_id_fk == person.id,
            Attendance.date >= from_date,
            Attendance.date <= to_date
        )
        .all()
    )

    attendance_map = {
        attendance.date: attendance
        for attendance in attendances
    }

    holidays = db.query(Holiday).all()

    report_rows = []

    current_date = from_date

    while current_date <= to_date:

        attendance = attendance_map.get(
            current_date
        )

        status, in_time, out_time = get_date_status(
            current_date,
            attendance,
            holidays
        )

        report_rows.append(
            {
                "date": current_date,
                "day": current_date.strftime("%A"),
                "status": status,
                "in_time": in_time,
                "out_time": out_time
            }
        )

        current_date += timedelta(days=1)

    return report_rows


# ============================================================
# CREATE PERSON-WISE EXCEL
# ============================================================

def create_person_excel(
    person,
    report_rows,
    from_date,
    to_date
):

    workbook = Workbook()

    ws = workbook.active

    ws.title = "Attendance Report"

    # --------------------------------------------------------
    # STUDENT INFORMATION
    # --------------------------------------------------------

    ws.append([
        "Student Attendance Report"
    ])

    ws.merge_cells(
        "A1:H1"
    )

    ws["A1"].font = Font(
        bold=True,
        size=16
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    ws.append([])

    ws.append([
        "Student ID",
        person.person_id
    ])

    ws.append([
        "Student Name",
        person.full_name
    ])

    ws.append([
        "Department",
        person.department
    ])

    ws.append([
        "Email",
        person.email
    ])

    ws.append([
        "Contact",
        person.contact
    ])

    ws.append([
        "Report From",
        from_date.strftime("%d-%m-%Y")
    ])

    ws.append([
        "Report To",
        to_date.strftime("%d-%m-%Y")
    ])

    ws.append([])

    # --------------------------------------------------------
    # CALCULATE SUMMARY
    # --------------------------------------------------------

    present_count = 0
    absent_count = 0
    sunday_count = 0
    holiday_count = 0

    for row in report_rows:

        status = row["status"]

        if status == "Present":

            present_count += 1

        elif status == "Absent":

            absent_count += 1

        elif status == "Sunday Holiday":

            sunday_count += 1

        else:

            holiday_count += 1

    working_days = (
        present_count +
        absent_count
    )

    if working_days > 0:

        attendance_percentage = (
            present_count /
            working_days
        ) * 100

    else:

        attendance_percentage = 0

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    ws.append([
        "Summary"
    ])

    ws.append([
        "Present",
        present_count
    ])

    ws.append([
        "Absent",
        absent_count
    ])

    ws.append([
        "Sunday Holidays",
        sunday_count
    ])

    ws.append([
        "Other Holidays",
        holiday_count
    ])

    ws.append([
        "Working Days",
        working_days
    ])

    ws.append([
        "Attendance Percentage",
        f"{attendance_percentage:.2f}%"
    ])

    ws.append([])

    # --------------------------------------------------------
    # ATTENDANCE TABLE
    # --------------------------------------------------------

    header_row = ws.max_row + 1

    headers = [
        "Date",
        "Day",
        "Status",
        "In Time",
        "Out Time"
    ]

    ws.append(headers)

    for row in report_rows:

        in_time = row["in_time"]

        out_time = row["out_time"]

        ws.append([
            row["date"].strftime("%d-%m-%Y"),
            row["day"],
            row["status"],
            in_time.strftime("%H:%M:%S")
            if in_time else "-",
            out_time.strftime("%H:%M:%S")
            if out_time else "-"
        ])

    # --------------------------------------------------------
    # STYLE TABLE HEADER
    # --------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for cell in ws[header_row]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------------
    # BORDER / ALIGNMENT
    # --------------------------------------------------------

    for row in ws.iter_rows(
        min_row=header_row,
        max_row=ws.max_row
    ):

        for cell in row:

            cell.border = thin_border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    # --------------------------------------------------------
    # WIDTH
    # --------------------------------------------------------

    auto_width_columns(ws)

    # --------------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------------

    ws.freeze_panes = f"A{header_row + 1}"

    # --------------------------------------------------------
    # SAVE TO MEMORY
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


# ============================================================
# PERSON-WISE EXCEL REPORT
# ============================================================

@app.get(
    "/reports/person-wise/{person_id}"
)
def download_person_report(
    person_id: str,

    from_date: Optional[date] = Query(
        default=None
    ),

    to_date: Optional[date] = Query(
        default=None
    ),

    db: Session = Depends(get_db)
):

    # --------------------------------------------------------
    # FIND PERSON
    # --------------------------------------------------------

    person = (
        db.query(Person)
        .filter(Person.person_id == person_id)
        .first()
    )

    # Support old numeric database IDs as well.
    if not person and person_id.isdigit():
        person = (
            db.query(Person)
            .filter(Person.id == int(person_id))
            .first()
        )

    if not person:

        raise HTTPException(
            status_code=404,
            detail="Person not found"
        )

    # --------------------------------------------------------
    # FULL REPORT
    # --------------------------------------------------------

    if from_date is None and to_date is None:

        from_date = person.created_at.date()

        to_date = datetime.now(IST).date()

    # --------------------------------------------------------
    # VALIDATE DATES
    # --------------------------------------------------------

    elif from_date is not None and to_date is not None:

        if to_date < from_date:

            raise HTTPException(
                status_code=400,
                detail=(
                    "To date cannot be before "
                    "from date"
                )
            )

    else:

        raise HTTPException(
            status_code=400,
            detail=(
                "Please provide both "
                "from_date and to_date"
            )
        )

    # --------------------------------------------------------
    # BUILD REPORT
    # --------------------------------------------------------

    report_rows = build_person_report_data(
        person,
        db,
        from_date,
        to_date
    )

    # --------------------------------------------------------
    # CREATE EXCEL
    # --------------------------------------------------------

    excel_file = create_person_excel(
        person,
        report_rows,
        from_date,
        to_date
    )

    filename = (
        f"{person.person_id}_{person.full_name}_"
        f"attendance_report.xlsx"
    )

    return StreamingResponse(
        excel_file,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )


# ============================================================
# DAY-WISE EXCEL REPORT
# ============================================================

@app.get(
    "/reports/day-wise"
)
def download_day_report(
    report_date: date = Query(...),
    db: Session = Depends(get_db)
):

    # --------------------------------------------------------
    # GET ALL PERSONS
    # --------------------------------------------------------

    persons = (
        db.query(Person)
        .order_by(Person.id.asc())
        .all()
    )

    # --------------------------------------------------------
    # GET ATTENDANCE FOR SELECTED DATE
    # --------------------------------------------------------

    attendances = (
        db.query(Attendance)
        .filter(
            Attendance.date == report_date
        )
        .all()
    )

    attendance_map = {
        attendance.person_id_fk: attendance
        for attendance in attendances
    }

    # --------------------------------------------------------
    # HOLIDAYS
    # --------------------------------------------------------

    holidays = db.query(Holiday).all()

    # --------------------------------------------------------
    # WORKBOOK
    # --------------------------------------------------------

    workbook = Workbook()

    ws = workbook.active

    ws.title = "Day Wise Attendance"

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    ws.append([
        "Day Wise Attendance Report"
    ])

    ws.merge_cells(
        "A1:H1"
    )

    ws["A1"].font = Font(
        bold=True,
        size=16
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    ws.append([])

    ws.append([
        "Report Date",
        report_date.strftime("%d-%m-%Y")
    ])

    ws.append([
        "Day",
        report_date.strftime("%A")
    ])

    ws.append([])

    # --------------------------------------------------------
    # BUILD ROW DATA
    # --------------------------------------------------------

    report_rows = []

    for person in persons:

        attendance = attendance_map.get(
            person.id
        )

        status, in_time, out_time = get_date_status(
            report_date,
            attendance,
            holidays
        )

        report_rows.append(
            {
                "person_id": person.person_id,
                "full_name": person.full_name,
                "department": person.department,
                "date": report_date,
                "day": report_date.strftime("%A"),
                "status": status,
                "in_time": in_time,
                "out_time": out_time
            }
        )

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    total_students = len(persons)

    present_count = sum(
        1
        for row in report_rows
        if row["status"] == "Present"
    )

    absent_count = sum(
        1
        for row in report_rows
        if row["status"] == "Absent"
    )

    sunday_count = sum(
        1
        for row in report_rows
        if row["status"] == "Sunday Holiday"
    )

    other_holiday_count = sum(
        1
        for row in report_rows
        if (
            row["status"] != "Present"
            and row["status"] != "Absent"
            and row["status"] != "Sunday Holiday"
        )
    )

    # --------------------------------------------------------
    # SUMMARY ROWS
    # --------------------------------------------------------

    ws.append([
        "Total Students",
        total_students
    ])

    ws.append([
        "Present",
        present_count
    ])

    ws.append([
        "Absent",
        absent_count
    ])

    ws.append([
        "Sunday Holiday",
        sunday_count
    ])

    ws.append([
        "Other Holiday",
        other_holiday_count
    ])

    ws.append([])

    # --------------------------------------------------------
    # TABLE HEADER
    # --------------------------------------------------------

    header_row = ws.max_row + 1

    headers = [
        "Person ID",
        "Student Name",
        "Department",
        "Date",
        "Day",
        "Status",
        "In Time",
        "Out Time"
    ]

    ws.append(headers)

    # --------------------------------------------------------
    # TABLE DATA
    # --------------------------------------------------------

    for row in report_rows:

        ws.append([
            row["person_id"],
            row["full_name"],
            row["department"],
            row["date"].strftime("%d-%m-%Y"),
            row["day"],
            row["status"],
            (
                row["in_time"].strftime("%H:%M:%S")
                if row["in_time"]
                else "-"
            ),
            (
                row["out_time"].strftime("%H:%M:%S")
                if row["out_time"]
                else "-"
            )
        ])

    # --------------------------------------------------------
    # STYLE
    # --------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for cell in ws[header_row]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center"
        )

    for row in ws.iter_rows(
        min_row=header_row,
        max_row=ws.max_row
    ):

        for cell in row:

            cell.border = thin_border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    auto_width_columns(ws)

    ws.freeze_panes = f"A{header_row + 1}"

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        f"day_wise_attendance_"
        f"{report_date.strftime('%Y-%m-%d')}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )